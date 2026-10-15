"""Build Demo 10's campaign artifacts and the harness's pristine copies.

Provenance only. NOT run during verification. Run with a python that has openpyxl
(the repo venv has it):

    .venv/Scripts/python.exe examples/claude-demo-kit/demos/10_autonomous_investigation/checks/_generate_repo_data.py

Writes, from one seeded model of campaign VAL-2026-09 (STORY.md, 2026-09-26, seed
20260926), so the dataset, the workbook and the logs agree with each other:

    input/data/validation_runs.csv     one row per run (LF; written by the Linux runner)
    input/data/validation_summary.xlsx campaign totals and breakdowns
    input/logs/estimator-worker.log    per-run gate activity
    input/logs/validation-runner.log   per-run scoring and the campaign verdict
    input/logs/scheduler.log           dispatch and infrastructure noise

and copies the editable text inputs (src/, tests/, config/) to checks/orig__*.txt so the
harness can resolve the agent's file:line citations against the files as they were
staged, even after the agent has fixed them.

Not generated here, copied from demo 7 byte for byte: input/docs/CVT-SRS-003.pdf and
input/data/gate_vectors.csv (the v2.4 bench capture). checks/measurement_gate_ref.m is
the same hidden reference gate as demo 7's.

The planted failure (STORY.md, incident of 2026-09-26): src/measurement_gate.m compares
d2 against an inlined 5.991 (chi-squared 95 %, m = 2) instead of the gamma it is
passed. cvt_model passes the configured per-mode value (planar 9.210, Doppler 13.277).
For Doppler streams (m = 4) d2 ~ chi2(4), so P(d2 > 5.991) = 0.20: about one good
measurement in five is dropped and the track drifts past 0.5 m. Planar streams lose
5 % instead of 1 % and still pass. The logged rejected d2 values are drawn from the
chi-squared tail beyond 5.991 for the run's m, so every one of them exceeds 5.991 and
most Doppler ones are below 13.277.

Decoys: build (v2.3.2 planar baseline runs pass, and so do v2.4.1 planar runs), worker
(estimator-02 is 15-20 % slower on runtime, inside the 50 ms limit, unrelated to
error), the v2.3.0 covariance regression (fixed in v2.3.2; kalman_update forms S from
the current P every step) and the Doppler noise/limit theory in the handover notes.
"""

from __future__ import annotations

import math
import random
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
IN = HERE.parent / "input"

SEED = 20260926
MEASUREMENTS = 1500
ERROR_LIMIT_M = 0.5            # CVT-VAL-002 §4
RUNTIME_LIMIT_MS = 50.0        # CVT-VAL-002 §4
MIN_PASS_RATE = 0.95
WORKERS = ["estimator-01", "estimator-02", "estimator-03"]
START = datetime(2026, 9, 26, 4, 3, 11, tzinfo=timezone.utc)
GATE_INLINED = 5.991
P_REJECT = {2: math.exp(-GATE_INLINED / 2),                                  # chi2(2)
            4: math.exp(-GATE_INLINED / 2) * (1 + GATE_INLINED / 2)}         # chi2(4)
ALERT_PCT = 10.0               # estimator_config.json logging.gate_rejection_alert_pct


def build_runs() -> list[dict]:
    rng = random.Random(SEED)
    kinds = ([("planar", 2, "v2.3.2")] * 20 + [("planar", 2, "v2.4.1")] * 20
             + [("doppler", 4, "v2.4.1")] * 20)
    rng.shuffle(kinds)
    doppler_idx = [i for i, k in enumerate(kinds) if k[0] == "doppler"]
    doppler_pass = set(rng.sample(doppler_idx, 3))

    runs, t = [], START
    for i, (mode, m, build) in enumerate(kinds):
        worker = WORKERS[i % 3]
        if mode == "planar":
            error = rng.uniform(0.07, 0.41)
        elif i in doppler_pass:
            error = rng.uniform(0.41, 0.48)
        else:
            error = rng.uniform(0.56, 1.38)
        runtime = rng.uniform(26.0, 34.0) * (rng.uniform(1.15, 1.20) if worker == "estimator-02" else 1.0)
        rejected = sum(1 for _ in range(MEASUREMENTS) if rng.random() < P_REJECT[m])
        status = "PASS" if (error <= ERROR_LIMIT_M and runtime <= RUNTIME_LIMIT_MS) else "FAIL"
        runs.append({
            "run_id": "VR-%d" % (2101 + i), "mode": mode, "measurement_dim": m,
            "build": build, "worker": worker, "max_pos_error_m": round(error, 4),
            "runtime_ms": round(runtime, 1), "status": status, "rejected": rejected,
            "reject_pct": 100.0 * rejected / MEASUREMENTS, "t_start": t,
        })
        t += timedelta(seconds=rng.randint(95, 160))
    return runs


def write_text(path: Path, lines: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(lines) + "\n")


def write_csv(runs: list[dict]) -> None:
    cols = ["run_id", "mode", "measurement_dim", "build", "worker",
            "max_pos_error_m", "runtime_ms", "status"]
    write_text(IN / "data" / "validation_runs.csv",
               [",".join(cols)] + [",".join(str(r[c]) for c in cols) for r in runs])


def write_workbook(runs: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    fixed = datetime(2026, 9, 26, 6, 15, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = "validation-runner"

    def summarize(key):
        groups: dict[str, list[dict]] = {}
        for r in runs:
            groups.setdefault(r[key], []).append(r)
        out = []
        for name in sorted(groups):
            g = groups[name]
            failed = sum(1 for r in g if r["status"] == "FAIL")
            errs = [r["max_pos_error_m"] for r in g]
            rts = [r["runtime_ms"] for r in g]
            out.append([name, len(g), len(g) - failed, failed,
                        round((len(g) - failed) / len(g), 4),
                        round(sum(errs) / len(errs), 4), round(max(errs), 4),
                        round(sum(rts) / len(rts), 1)])
        return out

    ov = wb.active
    ov.title = "Campaign"
    failed = sum(1 for r in runs if r["status"] == "FAIL")
    ov.append(["Metric", "Value"])
    for row in [
        ["Campaign", "VAL-2026-09"],
        ["Date", "2026-09-26"],
        ["Build under test", "v2.4.1"],
        ["Baseline build", "v2.3.2"],
        ["Total runs", len(runs)],
        ["Passed", len(runs) - failed],
        ["Failed", failed],
        ["Pass rate", round((len(runs) - failed) / len(runs), 4)],
        ["Min pass rate (CVT-VAL-002 §4)", MIN_PASS_RATE],
        ["Max position error limit (m)", ERROR_LIMIT_M],
        ["Max runtime limit (ms)", RUNTIME_LIMIT_MS],
        ["Verdict", "REJECT"],
    ]:
        ov.append(row)
    for c in ov[1]:
        c.font = Font(bold=True)
    ov.column_dimensions["A"].width = 32
    ov.column_dimensions["B"].width = 14

    header = ["Group", "Runs", "Passed", "Failed", "Pass rate",
              "Mean error (m)", "Max error (m)", "Mean runtime (ms)"]
    for title, key in [("By Mode", "mode"), ("By Build", "build"), ("By Worker", "worker")]:
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in summarize(key):
            ws.append(row)
        ws.column_dimensions["A"].width = 16
    out = IN / "data" / "validation_summary.xlsx"
    wb.save(out)
    # openpyxl stamps zip members with the current time; pin them so the file is
    # byte-for-byte reproducible.
    import zipfile
    with zipfile.ZipFile(out) as zin:
        members = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in members:
            if name == "docProps/core.xml":     # openpyxl sets modified = now on save
                import re
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-09-26T06:15:00Z", data)
            zi = zipfile.ZipInfo(name, date_time=(2026, 9, 26, 6, 15, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)


def _ts(t) -> str:
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


def rejected_d2(rng: random.Random, m: int) -> float:
    """A d2 the inlined gate rejects: chi-squared(m) conditioned on d2 > 5.991."""
    while True:
        d = rng.gammavariate(m / 2.0, 2.0)
        if d > GATE_INLINED:
            return d


def write_logs(runs: list[dict]) -> None:
    rng = random.Random(SEED + 1)
    worker = [
        "2026-09-26T04:02:58Z INFO  [supervisor] campaign VAL-2026-09 starting "
        "(build v2.4.1, baseline v2.3.2, 60 runs)",
        "2026-09-26T04:02:58Z INFO  [supervisor] config estimator_config.json version=2.4.1 "
        "gate.enabled=true gate.confidence=0.99",
        "2026-09-26T04:03:02Z INFO  [supervisor] 3 workers available: "
        "estimator-01, estimator-02, estimator-03",
    ]
    runner = [
        "2026-09-26T04:02:59Z INFO  [validation] campaign VAL-2026-09 queued: 60 runs",
        "2026-09-26T04:03:00Z INFO  [validation] acceptance: max_pos_error<=0.500 m "
        "runtime<=50.0 ms min_pass_rate>=0.950 (CVT-VAL-002 Rev A section 4)",
    ]
    for r in runs:
        t0, w, rid, mode = r["t_start"], r["worker"], r["run_id"], r["mode"]
        worker.append("%s INFO  [worker:%s] run=%s mode=%s build=%s starting "
                      "(measurements=%d, dim=%d)"
                      % (_ts(t0), w, rid, mode, r["build"], MEASUREMENTS, r["measurement_dim"]))
        ks = sorted(rng.sample(range(12, MEASUREMENTS - 5), 3))
        for j, k in enumerate(ks):
            worker.append("%s WARN  [gate] run=%s k=%d rejected d2=%.3f"
                          % (_ts(t0 + timedelta(seconds=3 + 8 * j)), rid, k,
                             rejected_d2(rng, r["measurement_dim"])))
        if r["reject_pct"] > ALERT_PCT:
            worker.append("%s WARN  [gate] run=%s rejection rate %.2f%% above alert level %.1f%%"
                          % (_ts(t0 + timedelta(seconds=27)), rid, r["reject_pct"], ALERT_PCT))
        worker.append("%s INFO  [gate] run=%s mode=%s complete: measurements=%d rejected=%d (%.2f%%)"
                      % (_ts(t0 + timedelta(seconds=30)), rid, mode, MEASUREMENTS,
                         r["rejected"], r["reject_pct"]))
        worker.append("%s INFO  [worker:%s] run=%s finished (%.1f ms)"
                      % (_ts(t0 + timedelta(seconds=31)), w, rid, r["runtime_ms"]))
        t1 = _ts(t0 + timedelta(seconds=33))
        if r["status"] == "PASS":
            runner.append("%s INFO  [validation] run=%s mode=%s build=%s max_pos_error=%.3f m "
                          "runtime=%.1f ms -> PASS"
                          % (t1, rid, mode, r["build"], r["max_pos_error_m"], r["runtime_ms"]))
        else:
            runner.append("%s ERROR [validation] run=%s mode=%s build=%s max_pos_error=%.3f m "
                          "runtime=%.1f ms -> FAIL (max_pos_error > 0.500 m)"
                          % (t1, rid, mode, r["build"], r["max_pos_error_m"], r["runtime_ms"]))

    failed = sum(1 for r in runs if r["status"] == "FAIL")
    end = runs[-1]["t_start"] + timedelta(seconds=75)
    runner.append("%s ERROR [validation] campaign VAL-2026-09 complete: %d runs, %d failed, "
                  "pass_rate=%.3f (min 0.950) -> REJECT"
                  % (_ts(end), len(runs), failed, (len(runs) - failed) / len(runs)))
    runner.append("%s INFO  [validation] summary written to validation_summary.xlsx"
                  % _ts(end + timedelta(seconds=2)))
    worker.append("%s INFO  [supervisor] campaign VAL-2026-09 complete" % _ts(end))

    # Scheduler: dispatch, the estimator-02 runtime flag and infrastructure noise.
    def mean(xs):
        return sum(xs) / len(xs)

    others = mean([r["runtime_ms"] for r in runs if r["worker"] != "estimator-02"])
    e02 = mean([r["runtime_ms"] for r in runs if r["worker"] == "estimator-02"])
    last = {w: max(r["t_start"] for r in runs if r["worker"] == w) + timedelta(seconds=32)
            for w in WORKERS}
    mid = runs[len(runs) // 2]["t_start"]
    sched = [
        (START - timedelta(seconds=31), "INFO  [scheduler] picked up campaign VAL-2026-09 (60 runs)"),
        (START - timedelta(seconds=30), "INFO  [scheduler] dispatch policy=round-robin workers=3"),
        (runs[7]["t_start"] + timedelta(seconds=12),
         "WARN  [scheduler] estimator-02 heartbeat late (1.8s) - continuing"),
        (runs[21]["t_start"] + timedelta(seconds=40),
         "WARN  [net] artifact upload retry 1/3 (connection reset by peer)"),
        (runs[21]["t_start"] + timedelta(seconds=42), "INFO  [net] artifact upload succeeded on retry"),
        (mid + timedelta(seconds=50), "WARN  [disk] /var/spool/cvt at 71% capacity"),
        (runs[44]["t_start"] + timedelta(seconds=61),
         "INFO  [cert] TLS certificate for artifact store expires in 34 days"),
    ]
    for w in WORKERS:
        sched.append((last[w], "INFO  [scheduler] %s completed its queue" % w))
    sched.append((end + timedelta(seconds=20),
                  "WARN  [scheduler] estimator-02 mean runtime %.0f%% above other workers - "
                  "flagged for review" % (100.0 * (e02 / others - 1.0))))
    sched.append((end + timedelta(seconds=21),
                  "INFO  [scheduler] campaign VAL-2026-09 released to validation"))
    sched.sort(key=lambda p: p[0])

    write_text(IN / "logs" / "estimator-worker.log", worker)
    write_text(IN / "logs" / "validation-runner.log", runner)
    write_text(IN / "logs" / "scheduler.log", ["%s %s" % (_ts(t), msg) for t, msg in sched])


def copy_pristine() -> None:
    for sub in ("src", "tests", "config"):
        for f in sorted((IN / sub).iterdir()):
            if f.is_file():
                shutil.copyfile(f, HERE / ("orig__%s__%s.txt" % (sub, f.name)))


def main() -> None:
    runs = build_runs()
    write_csv(runs)
    write_workbook(runs)
    write_logs(runs)
    copy_pristine()

    print("wrote %d runs" % len(runs))
    for key in ("mode", "build", "worker"):
        groups: dict[str, list[dict]] = {}
        for r in runs:
            groups.setdefault(r[key], []).append(r)
        for g in sorted(groups):
            rs = groups[g]
            f = sum(1 for r in rs if r["status"] == "FAIL")
            print("  %-7s %-13s n=%2d failed=%2d reject=%.2f%% runtime=%.1f ms"
                  % (key, g, len(rs), f, sum(r["reject_pct"] for r in rs) / len(rs),
                     sum(r["runtime_ms"] for r in rs) / len(rs)))
    print("  failing:", " ".join(r["run_id"] for r in runs if r["status"] == "FAIL"))


if __name__ == "__main__":
    main()
