"""Build reference/TraceabilityMatrix.xlsx - the golden traceability deliverable.

Provenance only (not run during verification). Run with a python that has openpyxl:

    python checks/_generate_reference_xlsx.py

The matrix audits the v2.2.0 kalman_update.m and test_kalman.m against CVT-EST-001
Rev C, in the columns the prompt fixes (Requirement, Implemented, Tests, Status,
Evidence) plus a few extra ones a reviewer would add.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "reference" / "TraceabilityMatrix.xlsx"

COLUMNS = [
    "Requirement", "Source", "Implemented", "Implementation", "Tests", "Result",
    "Status", "Evidence", "Action",
]

ROWS = [
    ["R1", "CVT-EST-001 Rev C §5 R1 (§3)", "Y", "kalman_update.m line 1: function X = kalman_update(Z, model)",
     "T1", "PASS", "VERIFIED",
     "T1 calls X = kalman_update(Z, model) and uses the result.", ""],
    ["R2", "CVT-EST-001 Rev C §5 R2 (§3)", "Y", "X = zeros(N, n); X(k, :) = x' per step, k = 1..N",
     "T1, T3", "PASS", "VERIFIED",
     "T1 asserts size(X) == [rows(Z), 4]; T3 compares X(:,1:2) with Z row by row, so order is exercised (loosely).",
     ""],
    ["R3", "CVT-EST-001 Rev C §5 R3 (§4)", "Y",
     "x = model.x0; P = model.P0; predict x = F*x, P = F*P*F' + Q; update K = (P*H')/S, x = x + K*y, P = (I - K*H)*P",
     "T3", "PASS", "PARTIAL",
     "T3 only bounds |X(:,1:2) - Z| < 5.0, which a badly wrong recursion could still pass, and every test uses "
     "x0 = 0, P0 = I, so initialization from model.x0/model.P0 is never exercised. No comparison against a "
     "reference computation.",
     "Add a test with non-default x0/P0 and a tight comparison against an independent evaluation of §4."],
    ["R4", "CVT-EST-001 Rev C §5 R4", "Y", "Pure function: no globals, files, or randomness",
     "T2", "PASS", "VERIFIED", "T2 asserts isequal(X, kalman_update(Z, model)).", ""],
    ["R5", "CVT-EST-001 Rev C §5 R5", "Y", "test_kalman.m (runs on test_vectors.csv)",
     "T1, T2, T3", "PASS (test_results.txt, 2026-03-10, v2.2.0)", "VERIFIED",
     "Suite is runnable and exercises R1-R4 on test_vectors.csv; 3 passed, 0 failed. Its R3 coverage is partial (see R3).",
     ""],
    ["R6", "CVT-EST-001 Rev C §5 R6", "N",
     "None: no isfinite/validity check on Z; a NaN row flows through y = z - H*x into x and P",
     "none", "-", "NOT_IMPLEMENTED",
     "No input validation in kalman_update.m and no test with NaN/Inf or malformed measurements.",
     "Implement rejection and reporting of non-finite or malformed measurements, and add tests for it."],
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Traceability"
ws.append(COLUMNS)
for c in ws[1]:
    c.font = Font(bold=True)
fills = {"VERIFIED": "C6EFCE", "PARTIAL": "FFEB9C", "NOT_VERIFIED": "FFEB9C", "NOT_IMPLEMENTED": "FFC7CE"}
for row in ROWS:
    ws.append(row)
    ws.cell(ws.max_row, COLUMNS.index("Status") + 1).fill = PatternFill("solid", fgColor=fills[row[6]])
for col, width in zip("ABCDEFGHI", (12, 26, 12, 50, 12, 22, 18, 70, 50)):
    ws.column_dimensions[col].width = width
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.alignment = Alignment(wrap_text=True, vertical="top")

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT} ({len(ROWS)} requirements)")
