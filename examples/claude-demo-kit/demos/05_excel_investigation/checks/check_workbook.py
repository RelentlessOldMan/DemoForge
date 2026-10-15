"""check_workbook.py  DemoForge verification harness for demo 5 (acceptance workbook).

Opens the agent's validation_results_corrected.xlsx and evaluates the workbook's
OWN formulas (xlcalc.py) instead of trusting any number or string in it:

  criteria_match_plan   the Criteria limits equal CVT-VAL-002 Rev A §4
  run_data_intact       the 40 recorded runs (ID, scenario, error, runtime) unchanged
  formulas_live         every Result cell, Passed, Pass rate and Verdict is a formula
  result_mismatches     Result cells that disagree with an independent per-run
                        recomputation (error <= 0.5 m AND runtime <= 50 ms)
  summary_mismatches    Passed / Failed / Pass rate cells that disagree
  verdict_mismatches    Verdict cells that disagree (REJECT iff pass rate < 0.95)
  verdict_reject        the delivered workbook's Verdict evaluates to REJECT

The three mismatch counts are summed over the delivered data plus five hidden
variants of the run data (written into the Runs cells before evaluation): all runs
within limits; three runs over the runtime limit; one slow run plus values exactly
on both limits; exactly 38/40 passing; 37/40 passing. A hard-coded "REJECT", a fix
that only changes the limit cell, or a formula that ignores runtime fails at least
one variant.

findings.json is checked against the same recomputation: verdict, counts, pass
rate, the exact set of failed runs, and that the issues cite the limit cell and
only cells that exist in the workbook as delivered to the agent.

Writes metrics.json. Exits non-zero only if it cannot run at all.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

import xlcalc

WORKBOOK = "validation_results_corrected.xlsx"
EXPECTED = json.loads(Path("expected.json").read_text(encoding="utf-8"))
PLAN = EXPECTED["plan"]
RUNS = EXPECTED["runs"]


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def truth(runs: list[dict]) -> dict:
    res = {r["run_id"]: (r["error"] <= PLAN["max_pos_error_m"] and r["runtime"] <= PLAN["max_runtime_ms"])
           for r in runs}
    passed = sum(res.values())
    rate = passed / len(runs)
    return {"result": res, "passed": passed, "failed": len(runs) - passed, "rate": rate,
            "verdict": "ACCEPT" if rate >= PLAN["min_pass_rate"] else "REJECT"}


def variants() -> list[tuple[str, list[dict]]]:
    """The delivered data plus hidden variants; only error/runtime values change."""
    base = [dict(r) for r in RUNS]

    def clean():
        out = [dict(r) for r in base]
        for i, r in enumerate(out):
            if r["error"] > 0.45:
                r["error"] = round(0.12 + 0.007 * i, 3)
        return out

    v1 = clean()
    v2 = clean()
    for i in (3, 17, 30):
        v2[i]["runtime"] = 55.0 + i / 10
    v3 = clean()
    v3[8]["runtime"] = 61.2
    v3[12]["error"] = 0.5
    v3[21]["runtime"] = 50.0
    v4 = clean()
    v4[5]["error"] = 0.51
    v4[26]["error"] = 0.74
    v5 = clean()
    for i in (2, 14, 33):
        v5[i]["error"] = 0.6 + i / 100
    return [("delivered", base), ("all-within", v1), ("3-slow", v2),
            ("1-slow+boundaries", v3), ("38-of-40", v4), ("37-of-40", v5)]


# ------------------------------------------------------------- locating -----

def label_rows(ws):
    """(row, label text, [cells to the right]) for rows whose first text cell is a label."""
    for row in ws.iter_rows():
        cells = [c for c in row if c.value is not None]
        if not cells or not isinstance(cells[0].value, str) or cells[0].value.startswith("="):
            continue
        yield cells[0].row, cells[0].value.strip(), [c for c in row if c.column > cells[0].column and c.value is not None]


def find_criteria(wb, calc):
    pats = {"max_pos_error_m": r"position", "min_pass_rate": r"pass\s*rate", "max_runtime_ms": r"run\s*time"}
    for ws in sorted(wb.worksheets, key=lambda w: "criteria" not in w.title.lower()):
        found = {}
        for _, label, rest in label_rows(ws):
            for key, pat in pats.items():
                if key not in found and re.search(pat, label, re.I) and rest:
                    v = calc.value(ws.title, rest[0].coordinate)
                    if is_num(v):
                        found[key] = float(v)
        if len(found) == 3:
            return ws.title, found
    return None, {}


def find_runs(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=10):
            heads = {c.column: str(c.value).strip() for c in row if isinstance(c.value, str)}
            idc = next((k for k, h in heads.items() if re.fullmatch(r"run\s*_?id", h, re.I)), None)
            if idc is None:
                continue

            def col(pat, exclude=None):
                for k, h in heads.items():
                    if re.search(pat, h, re.I) and not (exclude and re.search(exclude, h, re.I)):
                        return k
                return None

            cols = {"id": idc, "scenario": col(r"scenario"), "error": col(r"err", r"run\s*time"),
                    "runtime": col(r"run\s*time", r"(ok|pass|check|flag)"), "result": col(r"^result", None)}
            if cols["result"] is None:
                cols["result"] = col(r"^(pass|status|outcome)", None)
            rows = {}
            for r in range(row[0].row + 1, ws.max_row + 1):
                v = ws.cell(row=r, column=idc).value
                if isinstance(v, str) and re.fullmatch(r"VR-\d+", v.strip()):
                    rows[v.strip()] = r
            return ws.title, cols, rows
    return None, {}, {}


def find_summary(wb, skip):
    want = {"verdict": r"^verdict", "passed": r"^(runs\s+)?passed\b", "failed": r"^(runs\s+)?failed\b",
            "rate": r"^pass\s*rate"}
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        cells = {}
        for _, label, rest in label_rows(ws):
            for key, pat in want.items():
                if key not in cells and re.search(pat, label, re.I) and rest:
                    cells[key] = rest[0].coordinate
        if "verdict" in cells:
            return ws.title, cells
    return None, {}


def as_result(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, str) and v.strip().upper() in ("PASS", "FAIL"):
        return v.strip().upper() == "PASS"
    return None


# ------------------------------------------------------------- findings -----

CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?$", re.I)


def split_cell(sheet, cell):
    """Accept {"sheet": "Criteria", "cell": "B2"} or a cell written as "Criteria!B2"."""
    sheet = str(sheet or "").strip().strip("'")
    cell = str(cell or "").strip()
    if "!" in cell:
        s, cell = cell.rsplit("!", 1)
        sheet = sheet or s.strip("'")
    return sheet, cell


def cell_in_layout(sheet, cell) -> bool:
    layout = {k.lower(): v for k, v in EXPECTED["input_layout"].items()}
    dims = layout.get(sheet.lower())
    m = CELL_RE.match(cell)
    if not dims or not m:
        return False
    for col, row in ((m.group(1), m.group(2)), (m.group(3), m.group(4))):
        if col is None:
            continue
        if int(row) < 1 or int(row) > dims[0] or column_index_from_string(col.upper()) > dims[1]:
            return False
    return True


def check_findings(t0: dict) -> dict:
    out = {"findings_valid": 0, "findings_verdict_reject": 0, "findings_counts_ok": 0,
           "findings_failed_runs_ok": 0, "findings_cites_limit_cell": 0, "findings_citations_valid": 0}
    p = Path("findings.json")
    if not p.exists():
        return out
    try:
        f = json.loads(p.read_text(encoding="utf-8"))
        verdict = str(f["campaign_verdict"]).strip().upper()
        total, passed, rate = f["runs_total"], f["runs_passed"], float(f["pass_rate"])
        failed = {str(x).strip().upper() for x in f["failed_runs"]}
        issues = f["issues"]
        assert isinstance(issues, list)
    except Exception as exc:  # noqa: BLE001 - any malformed findings is a failed check
        print(f"findings.json invalid: {exc}")
        return out
    out["findings_valid"] = 1
    out["findings_verdict_reject"] = int(verdict == "REJECT" == t0["verdict"])
    out["findings_counts_ok"] = int(total == len(RUNS) and passed == t0["passed"] and abs(rate - t0["rate"]) < 1e-3)
    true_failed = {k for k, ok in t0["result"].items() if not ok}
    out["findings_failed_runs_ok"] = int(failed == true_failed)
    lim_sheet, lim_cell = EXPECTED["limit_cell"]
    valid = bool(issues)
    for it in issues:
        if not isinstance(it, dict):
            valid = False
            continue
        sheet, cell = split_cell(it.get("sheet"), it.get("cell"))
        if not cell_in_layout(sheet, cell):
            valid = False
            continue
        m = CELL_RE.match(cell)
        if (sheet.lower() == lim_sheet.lower() and m.group(3) is None
                and f"{m.group(1).upper()}{m.group(2)}" == lim_cell):
            out["findings_cites_limit_cell"] = 1
    out["findings_citations_valid"] = int(valid)
    return out


# ----------------------------------------------------------------- main -----

def main() -> int:
    t0 = truth(RUNS)
    metrics = {"n_runs": len(RUNS), "recomputed_pass_rate": round(t0["rate"], 4),
               "n_failing_runs": t0["failed"], "workbook_present": 0, "criteria_match_plan": 0,
               "run_data_intact": 0, "formulas_live": 0, "result_mismatches": -1,
               "summary_mismatches": -1, "verdict_mismatches": -1, "verdict_reject": 0,
               "unsupported_formula": 0}
    metrics.update(check_findings(t0))

    if Path(WORKBOOK).exists():
        metrics["workbook_present"] = 1
        try:
            wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
            evaluate(wb, metrics)
        except xlcalc.Unsupported as exc:
            metrics["unsupported_formula"] = 1
            print(f"cannot evaluate the workbook's formulas: {exc}")
    else:
        print(f"{WORKBOOK} not found")

    Path("metrics.json").write_text(json.dumps(metrics), encoding="utf-8")
    print(json.dumps(metrics, indent=1))
    return 0


def evaluate(wb, metrics: dict) -> None:
    calc = xlcalc.Workbook(wb)
    crit_sheet, crit = find_criteria(wb, calc)
    metrics["criteria_match_plan"] = int(len(crit) == 3 and all(abs(crit[k] - PLAN[k]) < 1e-9 for k in PLAN))
    print(f"criteria on {crit_sheet!r}: {crit}")

    runs_sheet, cols, rows = find_runs(wb)
    if not runs_sheet or None in (cols.get("error"), cols.get("runtime"), cols.get("result")):
        print(f"could not locate the run table: sheet={runs_sheet!r} cols={cols}")
        return
    ws = wb[runs_sheet]
    intact = set(rows) == {r["run_id"] for r in RUNS} and len(rows) == len(RUNS)
    for r in RUNS:
        if not intact:
            break
        k = rows[r["run_id"]]
        err = calc.value(runs_sheet, f"{get_column_letter(cols['error'])}{k}")
        rt = calc.value(runs_sheet, f"{get_column_letter(cols['runtime'])}{k}")
        scen = ws.cell(row=k, column=cols["scenario"]).value if cols.get("scenario") else r["scenario"]
        intact = (is_num(err) and abs(err - r["error"]) < 1e-9 and is_num(rt) and abs(rt - r["runtime"]) < 1e-9
                  and str(scen).strip() == r["scenario"])
    metrics["run_data_intact"] = int(intact)
    if not intact:
        print("the recorded run data was changed")
        return

    sum_sheet, sumc = find_summary(wb, {crit_sheet, runs_sheet})
    if not sum_sheet or "passed" not in sumc or "rate" not in sumc:
        print(f"could not locate Summary cells: {sum_sheet!r} {sumc}")
        return
    res_col = get_column_letter(cols["result"])
    formula = lambda sh, c: xlcalc.Workbook.formula_text(wb[sh][c].value) is not None  # noqa: E731
    live = all(formula(runs_sheet, f"{res_col}{k}") for k in rows.values())
    live = live and all(formula(sum_sheet, sumc[k]) for k in ("verdict", "passed", "rate"))
    metrics["formulas_live"] = int(live)

    res_mis = sum_mis = ver_mis = 0
    for name, data in variants():
        t = truth(data)
        over = {}
        for r in data:
            k = rows[r["run_id"]]
            over[(runs_sheet, f"{get_column_letter(cols['error'])}{k}")] = r["error"]
            over[(runs_sheet, f"{get_column_letter(cols['runtime'])}{k}")] = r["runtime"]
        c = xlcalc.Workbook(wb, over)
        bad_rows = [r["run_id"] for r in data
                    if as_result(c.value(runs_sheet, f"{res_col}{rows[r['run_id']]}")) is not t["result"][r["run_id"]]]
        passed = c.value(sum_sheet, sumc["passed"])
        rate = c.value(sum_sheet, sumc["rate"])
        verdict = c.value(sum_sheet, sumc["verdict"])
        s_bad = int(not (is_num(passed) and passed == t["passed"]))
        s_bad += int(not (is_num(rate) and abs(rate - t["rate"]) < 1e-3))
        if "failed" in sumc:
            failed = c.value(sum_sheet, sumc["failed"])
            s_bad += int(not (is_num(failed) and failed == t["failed"]))
        v_bad = int(not (isinstance(verdict, str) and verdict.strip().upper() == t["verdict"]))
        res_mis += len(bad_rows)
        sum_mis += s_bad
        ver_mis += v_bad
        print(f"  {name:18s} expect {t['passed']}/{len(data)} {t['verdict']:6s} | workbook "
              f"passed={passed} rate={rate} verdict={verdict} | bad results={bad_rows[:4]}"
              f"{'...' if len(bad_rows) > 4 else ''}")
        if name == "delivered":
            metrics["verdict_reject"] = int(isinstance(verdict, str) and verdict.strip().upper() == "REJECT")
            metrics["workbook_pass_rate"] = round(rate, 4) if is_num(rate) else None
    metrics["result_mismatches"] = res_mis
    metrics["summary_mismatches"] = sum_mis
    metrics["verdict_mismatches"] = ver_mis


if __name__ == "__main__":
    sys.exit(main())
