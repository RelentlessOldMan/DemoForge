"""Build the demo 5 acceptance workbook (input), its corrected reference, and the
harness's hidden expectations. Provenance only; not run during verification.

    .venv\\Scripts\\python.exe checks\\_generate_workbook.py

Story (STORY.md, demo 5): campaign VAL-2026-05, 2026-05-20 to 05-27, acceptance of
the v2.3.0 release candidate against CVT-VAL-002 Rev A (2026-05-12). Every run is
the same build; the scenario is what varies.

input/validation_results.xlsx
    Criteria, Runs and Summary sheets with live cross-sheet formulas and cached
    values, as Excel would have saved them. Two defects:
      1. Criteria!B2, the per-run position-error limit, reads 5 instead of 0.5. The
         six `maneuver` runs (0.6 to 1.1 m) therefore pass and the Summary shows
         ACCEPT; the correct outcome is REJECT (34/40 = 0.85 < 0.95).
      2. The per-run Result formula tests only the position error. The plan (§3)
         also requires runtime <= 50 ms per run. No run in this campaign is slow,
         so it changes nothing today, but the workbook is the campaign template
         and does not implement the plan. The harness tests it with hidden data.
    Run VR-1923 sits exactly on the 0.5 m limit, which passes (<=).
    estimator-02 runs 15-20 % slower than estimator-01 (the standing red herring).

reference/validation_results_corrected.xlsx
    Both defects fixed, formulas kept.

checks/expected.json
    Run data, plan limits and the input workbook's layout for the harness.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import random
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
DEMO = HERE.parent
sys.path.insert(0, str(HERE))
import xlcalc  # noqa: E402

INPUT_XLSX = DEMO / "input" / "validation_results.xlsx"
REF_XLSX = DEMO / "reference" / "validation_results_corrected.xlsx"
EXPECTED = HERE / "expected.json"

SEED = 20260520
PLAN = {"max_pos_error_m": 0.5, "min_pass_rate": 0.95, "max_runtime_ms": 50.0}
DEFECT_LIMIT = 5
BUILD = "v2.3.0-rc1"
FIRST_RUN = 1901
SCENARIO_MIX = {"cruise": 10, "turn": 9, "crossing": 8, "stop-go": 7, "maneuver": 6}
DAYS = [dt.date(2026, 5, 20) + dt.timedelta(days=d) for d in range(8)]


def build_runs() -> list[dict]:
    rng = random.Random(SEED)
    scen = [s for s, n in SCENARIO_MIX.items() for _ in range(n)]
    rng.shuffle(scen)
    runs = []
    for i, s in enumerate(scen):
        rid = FIRST_RUN + i
        worker = "estimator-01" if rid % 2 else "estimator-02"
        if s == "maneuver":
            err = rng.uniform(0.62, 1.08)
        elif s == "cruise":
            err = rng.uniform(0.09, 0.24)
        elif s == "crossing":
            err = rng.uniform(0.14, 0.33)
        else:  # turn, stop-go
            err = rng.uniform(0.21, 0.46)
        rt = rng.uniform(30.0, 37.0)
        if worker == "estimator-02":
            rt *= rng.uniform(1.15, 1.20)
        runs.append({
            "run_id": f"VR-{rid}",
            "date": DAYS[i // 5].isoformat(),
            "scenario": s,
            "build": BUILD,
            "worker": worker,
            "error": round(err, 3),
            "runtime": round(rt, 1),
        })
    # The turn run with the largest error is put exactly on the limit (it passes).
    b = max((r for r in runs if r["scenario"] == "turn"), key=lambda r: r["error"])
    b["error"] = 0.5
    return runs


def write_workbook(path: Path, runs: list[dict], corrected: bool) -> None:
    wb = openpyxl.Workbook()
    wb.properties.creator = "D. Volkov"
    wb.properties.lastModifiedBy = "D. Volkov"
    wb.properties.created = wb.properties.modified = dt.datetime(2026, 5, 28, 9, 0, 0)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9E1F2")

    cr = wb.active
    cr.title = "Criteria"
    cr.append(["Parameter", "Value", "Unit", "Source"])
    cr.append(["Max peak position error (per run)", 0.5 if corrected else DEFECT_LIMIT, "m", "CVT-VAL-002 Rev A §4"])
    cr.append(["Min pass rate", PLAN["min_pass_rate"], "fraction", "CVT-VAL-002 Rev A §4"])
    cr.append(["Max runtime (per run)", int(PLAN["max_runtime_ms"]), "ms", "CVT-VAL-002 Rev A §4"])
    for col, w in zip("ABCD", (36, 10, 10, 24)):
        cr.column_dimensions[col].width = w

    rn = wb.create_sheet("Runs")
    rn.append(["RunID", "Date", "Scenario", "Build", "Worker", "PeakPosError_m", "Runtime_ms", "Result"])
    for k, r in enumerate(runs, start=2):
        rn.append([r["run_id"], dt.date.fromisoformat(r["date"]), r["scenario"], r["build"],
                   r["worker"], r["error"], r["runtime"], None])
        rn.cell(row=k, column=2).number_format = "yyyy-mm-dd"
        rn.cell(row=k, column=6).number_format = "0.000"
        rn.cell(row=k, column=7).number_format = "0.0"
        if corrected:
            f = f'=IF(AND(F{k}<=Criteria!$B$2,G{k}<=Criteria!$B$4),"PASS","FAIL")'
        else:
            f = f'=IF(F{k}<=Criteria!$B$2,"PASS","FAIL")'
        rn.cell(row=k, column=8).value = f
    last = len(runs) + 1
    for col, w in zip("ABCDEFGH", (10, 12, 11, 12, 14, 16, 12, 9)):
        rn.column_dimensions[col].width = w
    rn.freeze_panes = "A2"

    sm = wb.create_sheet("Summary")
    rows = [
        ["Item", "Value"],
        ["Campaign", "VAL-2026-05"],
        ["Build under test", BUILD],
        ["Period", f"{DAYS[0].isoformat()} to {DAYS[-1].isoformat()}"],
        ["Total runs", f"=COUNTA(Runs!A2:A{last})"],
        ["Passed", f'=COUNTIF(Runs!H2:H{last},"PASS")'],
        ["Failed", f'=COUNTIF(Runs!H2:H{last},"FAIL")'],
        ["Pass rate", "=B6/B5"],
        ["Mean peak position error (m)", f"=AVERAGE(Runs!F2:F{last})"],
        ["Max peak position error (m)", f"=MAX(Runs!F2:F{last})"],
        ["Max runtime (ms)", f"=MAX(Runs!G2:G{last})"],
        ["Verdict", '=IF(B8>=Criteria!B3,"ACCEPT","REJECT")'],
        ["Prepared by", "D. Volkov, 2026-05-28"],
    ]
    for row in rows:
        sm.append(row)
    sm["B8"].number_format = "0.0%"
    sm["B9"].number_format = "0.000"
    sm["B10"].number_format = "0.000"
    sm["B11"].number_format = "0.0"
    sm["A12"].font = bold
    sm["B12"].font = bold
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 24

    for ws in (cr, rn, sm):
        for c in ws[1]:
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_with_cached_values(buf.getvalue()))


def _with_cached_values(xlsx: bytes) -> bytes:
    """Fill in each formula's cached result, as Excel does when it saves.

    openpyxl writes formulas without results, so a reader using data_only=True would
    see blanks. The values come from evaluating the workbook's own formulas.
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    calc = xlcalc.Workbook(wb)
    sheet_xml = {f"xl/worksheets/sheet{i}.xml": ws.title for i, ws in enumerate(wb.worksheets, start=1)}
    cell_re = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>')

    def fill(title):
        def repl(m):
            coord, attrs, formula = m.group(1), m.group(2), m.group(3)
            val = calc.value(title, coord)
            attrs = re.sub(r'\s*t="[^"]*"', "", attrs)
            if isinstance(val, str):
                return f'<c r="{coord}"{attrs} t="str"><f>{formula}</f><v>{escape(val)}</v></c>'
            if isinstance(val, bool):
                return f'<c r="{coord}"{attrs} t="b"><f>{formula}</f><v>{int(val)}</v></c>'
            return f'<c r="{coord}"{attrs}><f>{formula}</f><v>{repr(float(val))}</v></c>'
        return repl

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in sheet_xml:
                text = data.decode("utf-8")
                text, n = cell_re.subn(fill(sheet_xml[item.filename]), text)
                data = text.encode("utf-8")
            item.date_time = (2026, 5, 28, 9, 0, 0)  # stable bytes across regenerations
            zout.writestr(item, data)
    return out.getvalue()


def main() -> None:
    runs = build_runs()
    write_workbook(INPUT_XLSX, runs, corrected=False)
    write_workbook(REF_XLSX, runs, corrected=True)

    wb = openpyxl.load_workbook(INPUT_XLSX)
    layout = {ws.title: [ws.max_row, ws.max_column] for ws in wb.worksheets}
    EXPECTED.write_text(json.dumps({
        "plan": PLAN,
        "limit_cell": ["Criteria", "B2"],
        "input_layout": layout,
        "runs": runs,
    }, indent=1) + "\n", encoding="utf-8", newline="\n")

    passed = [r for r in runs if r["error"] <= PLAN["max_pos_error_m"] and r["runtime"] <= PLAN["max_runtime_ms"]]
    rate = len(passed) / len(runs)
    errs = sorted(r["error"] for r in runs if r["scenario"] == "maneuver")
    print(f"runs={len(runs)} passed={len(passed)} rate={rate:.3f} "
          f"verdict={'ACCEPT' if rate >= PLAN['min_pass_rate'] else 'REJECT'}")
    print(f"maneuver errors {errs}")
    print("max runtime by worker:", {w: max(r['runtime'] for r in runs if r['worker'] == w)
                                     for w in ('estimator-01', 'estimator-02')})
    calc = xlcalc.Workbook(wb)
    print("input workbook says:", calc.value("Summary", "B6"), calc.value("Summary", "B8"), calc.value("Summary", "B12"))


if __name__ == "__main__":
    main()
