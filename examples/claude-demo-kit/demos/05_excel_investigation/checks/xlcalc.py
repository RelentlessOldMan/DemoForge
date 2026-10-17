"""xlcalc.py  A small Excel formula evaluator for the demo 5 harness.

openpyxl reads formulas but cannot calculate them, and a workbook saved by openpyxl
carries no cached values. The harness therefore evaluates the workbook's own
formulas itself, optionally with some input cells overridden, to check that the
formulas (not just the numbers on screen) implement the validation plan.

Covers what a validation workbook realistically uses: cell and range references
(also cross-sheet and quoted sheet names, whole columns, defined names), arithmetic,
comparison and `&` operators with array broadcasting, IF/IFS/IFERROR/IFNA/LET,
and the functions in FUNCTIONS below (aggregates, the *IF/*IFS family, INDEX, MATCH,
VLOOKUP, XLOOKUP, ...). `_selftest_xlcalc.py` exercises each of them. An unsupported function raises Unsupported, which the harness
reports instead of guessing.
"""

from __future__ import annotations

import fnmatch
import math
import re

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter


class Unsupported(Exception):
    """The formula uses something this evaluator does not implement."""


class XLError:
    def __init__(self, code: str):
        self.code = code

    def __repr__(self) -> str:
        return self.code

    def __eq__(self, other) -> bool:
        return isinstance(other, XLError) and other.code == self.code

    def __hash__(self) -> int:
        return hash(self.code)


class _Raise(Exception):
    def __init__(self, err: XLError):
        self.err = err


class Array:
    """A 2-D block of values (a range or an array expression)."""

    def __init__(self, rows: list[list]):
        self.rows = rows

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.rows), (len(self.rows[0]) if self.rows else 0)

    def flat(self) -> list:
        return [v for row in self.rows for v in row]


# ---------------------------------------------------------------- parsing ----

_PREC = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}


def parse(formula: str):
    tokens = [t for t in Tokenizer(formula).items if t.type != Token.WSPACE]
    pos = 0

    def peek():
        return tokens[pos] if pos < len(tokens) else None

    def take():
        nonlocal pos
        t = tokens[pos]
        pos += 1
        return t

    def primary():
        t = take()
        if t.type == Token.OP_PRE:
            node = ("neg", unary_operand()) if t.value == "-" else unary_operand()
            return node
        if t.type == Token.OPERAND:
            if t.subtype == Token.NUMBER:
                return ("lit", float(t.value))
            if t.subtype == Token.TEXT:
                return ("lit", t.value[1:-1].replace('""', '"'))
            if t.subtype == Token.LOGICAL:
                return ("lit", t.value.upper() == "TRUE")
            if t.subtype == Token.ERROR:
                return ("lit", XLError(t.value))
            ref = t.value[6:] if t.value.lower().startswith("_xlpm.") else t.value  # LET names
            return ("ref", ref)
        if t.type == Token.PAREN and t.subtype == Token.OPEN:
            node = expr(0)
            close = take()
            if close.type != Token.PAREN:
                raise Unsupported(f"unbalanced parenthesis in {formula!r}")
            return node
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            name = t.value[:-1].upper()
            for prefix in ("_XLFN._XLWS.", "_XLFN.", "_XLWS."):
                if name.startswith(prefix):
                    name = name[len(prefix):]
            args = []
            if peek() is not None and peek().type == Token.FUNC and peek().subtype == Token.CLOSE:
                take()
                return ("func", name, args)
            while True:
                if peek() is not None and peek().type == Token.SEP:
                    args.append(("lit", None))  # empty argument
                else:
                    args.append(expr(0))
                sep = take()
                if sep.type == Token.FUNC and sep.subtype == Token.CLOSE:
                    return ("func", name, args)
                if sep.type != Token.SEP or sep.subtype != Token.ARG:
                    raise Unsupported(f"unexpected token {sep.value!r} in {formula!r}")
        raise Unsupported(f"unexpected token {t.value!r} in {formula!r}")

    def unary_operand():
        node = primary()
        while peek() is not None and peek().type == Token.OP_POST:
            take()
            node = ("pct", node)
        return node

    def expr(min_prec):
        left = unary_operand()
        while True:
            t = peek()
            if t is None or t.type != Token.OP_IN or t.value not in _PREC:
                if t is not None and t.type == Token.OP_IN:
                    raise Unsupported(f"operator {t.value!r} in {formula!r}")
                return left
            prec = _PREC[t.value]
            if prec < min_prec:
                return left
            take()
            right = expr(prec + (0 if t.value == "^" else 1))
            left = ("bin", t.value, left, right)

    node = expr(0)
    if pos != len(tokens):
        raise Unsupported(f"trailing tokens in {formula!r}")
    return node


# ------------------------------------------------------------- evaluation ----

_REF_RE = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<a>\$?[A-Z]{1,3}\$?\d*)(?::(?P<b>\$?[A-Z]{1,3}\$?\d*))?$", re.I)


def _num(v):
    if isinstance(v, XLError):
        raise _Raise(v)
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            raise _Raise(XLError("#VALUE!"))
    raise _Raise(XLError("#VALUE!"))


def _text(v) -> str:
    if isinstance(v, XLError):
        raise _Raise(v)
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _bool(v) -> bool:
    if isinstance(v, XLError):
        raise _Raise(v)
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    if isinstance(v, (int, float)):
        return v != 0
    if isinstance(v, str) and v.upper() in ("TRUE", "FALSE"):
        return v.upper() == "TRUE"
    raise _Raise(XLError("#VALUE!"))


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _compare(op: str, a, b) -> bool:
    for v in (a, b):
        if isinstance(v, XLError):
            raise _Raise(v)
    if a is None:
        a = 0.0 if _is_num(b) else ("" if isinstance(b, str) else False if isinstance(b, bool) else 0.0)
    if b is None:
        b = 0.0 if _is_num(a) else ("" if isinstance(a, str) else False if isinstance(a, bool) else 0.0)

    def rank(v):
        return 0 if _is_num(v) else 1 if isinstance(v, str) else 2

    if rank(a) != rank(b):
        ka, kb = rank(a), rank(b)
    elif isinstance(a, str):
        ka, kb = a.lower(), b.lower()
    else:
        ka, kb = float(a), float(b)
    return {"=": ka == kb, "<>": ka != kb, "<": ka < kb, ">": ka > kb, "<=": ka <= kb, ">=": ka >= kb}[op]


def _binop(op: str, a, b):
    if op in ("=", "<>", "<", ">", "<=", ">="):
        return _compare(op, a, b)
    if op == "&":
        return _text(a) + _text(b)
    x, y = _num(a), _num(b)
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if op == "/":
        if y == 0:
            raise _Raise(XLError("#DIV/0!"))
        return x / y
    if op == "^":
        return x ** y
    raise Unsupported(op)


def _elementwise(fn, a, b):
    """Apply a scalar binary function with Excel-style array broadcasting."""
    if not isinstance(a, Array) and not isinstance(b, Array):
        return fn(a, b)
    ra = a.rows if isinstance(a, Array) else None
    rb = b.rows if isinstance(b, Array) else None
    nr = max(len(ra) if ra else 1, len(rb) if rb else 1)
    nc = max(len(ra[0]) if ra else 1, len(rb[0]) if rb else 1)

    def at(m, src, i, j):
        if src is None:
            return m
        ii = 0 if len(src) == 1 else i
        jj = 0 if len(src[0]) == 1 else j
        if ii >= len(src) or jj >= len(src[0]):
            return XLError("#N/A")
        return src[ii][jj]

    out = []
    for i in range(nr):
        row = []
        for j in range(nc):
            try:
                row.append(fn(at(a, ra, i, j), at(b, rb, i, j)))
            except _Raise as e:
                row.append(e.err)
        out.append(row)
    return Array(out)


def _criterion(crit):
    """COUNTIF-style criterion -> predicate."""
    if isinstance(crit, Array):
        raise Unsupported("array criterion")
    if _is_num(crit) or isinstance(crit, bool):
        return lambda v: (_is_num(v) and float(v) == float(crit)) if _is_num(crit) else v is crit
    s = _text(crit)
    m = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", s, re.S)
    op, operand = m.group(1) or "=", m.group(2)
    try:
        num = float(operand)
    except ValueError:
        num = None
    if num is not None:
        def pred(v):
            if _is_num(v):
                return _compare(op, float(v), num)
            if isinstance(v, str):
                try:
                    return _compare(op, float(v), num)
                except ValueError:
                    return op == "<>"
            return op == "<>"
        return pred
    if operand == "":
        return (lambda v: v is None or v == "") if op == "=" else (lambda v: not (v is None or v == ""))
    pattern = operand.lower()
    wild = any(c in pattern for c in "*?")

    def spred(v):
        if not isinstance(v, str):
            hit = False
        elif wild:
            hit = fnmatch.fnmatchcase(v.lower(), pattern)
        else:
            hit = v.lower() == pattern
        if op == "=":
            return hit
        if op == "<>":
            return not hit
        return isinstance(v, str) and _compare(op, v, operand)
    return spred


def _vals(args, numbers_only=False):
    """Flatten aggregate arguments: ranges skip text/bools, direct args are coerced."""
    out = []
    for a in args:
        if isinstance(a, Array):
            for v in a.flat():
                if isinstance(v, XLError):
                    raise _Raise(v)
                if _is_num(v):
                    out.append(float(v))
        else:
            if a is None:
                continue
            out.append(_num(a))
    return out


def _ifs_mask(args):
    if len(args) % 2:
        raise _Raise(XLError("#VALUE!"))
    mask = None
    for rng, crit in zip(args[::2], args[1::2]):
        if not isinstance(rng, Array):
            raise _Raise(XLError("#VALUE!"))
        pred = _criterion(crit)
        cur = [pred(v) for v in rng.flat()]
        mask = cur if mask is None else [x and y for x, y in zip(mask, cur)]
    return mask


class Workbook:
    """Evaluates a loaded openpyxl workbook (data_only=False) with overrides."""

    def __init__(self, wb, overrides: dict | None = None):
        self.wb = wb
        self.sheets = {ws.title.lower(): ws for ws in wb.worksheets}
        self.overrides = {(s.lower(), c.upper().replace("$", "")): v for (s, c), v in (overrides or {}).items()}
        self.memo: dict = {}
        self.stack: set = set()
        self.scopes: list[dict] = []
        self.names = {}
        try:
            for name, dn in wb.defined_names.items():
                self.names[name.lower()] = dn.attr_text
        except AttributeError:  # older openpyxl
            for dn in getattr(wb.defined_names, "definedName", []):
                self.names[dn.name.lower()] = dn.attr_text

    # -- public --------------------------------------------------------------
    def value(self, sheet: str, coord: str):
        """The evaluated value of one cell (an XLError for an error result)."""
        key = (sheet.lower(), coord.upper().replace("$", ""))
        if key in self.overrides:
            return self.overrides[key]
        if key in self.memo:
            return self.memo[key]
        ws = self.sheets.get(key[0])
        if ws is None:
            return XLError("#REF!")
        raw = ws[key[1]].value
        formula = self.formula_text(raw)
        if formula is None:
            val = raw
            if isinstance(val, int) and not isinstance(val, bool):
                val = float(val)
        else:
            if key in self.stack:
                raise Unsupported(f"circular reference at {sheet}!{coord}")
            self.stack.add(key)
            try:
                val = self._eval(parse(formula), key[0])
                if isinstance(val, Array):  # implicit intersection: take the first cell
                    val = val.rows[0][0] if val.rows and val.rows[0] else None
            except _Raise as e:
                val = e.err
            finally:
                self.stack.discard(key)
        self.memo[key] = val
        return val

    @staticmethod
    def formula_text(raw):
        if isinstance(raw, str) and raw.startswith("=") and len(raw) > 1:
            return raw
        text = getattr(raw, "text", None)  # openpyxl ArrayFormula / DataTableFormula
        if isinstance(text, str):
            return text if text.startswith("=") else "=" + text
        return None

    # -- internals -----------------------------------------------------------
    def _range(self, ref: str, cur_sheet: str):
        if ref.lower() in self.names:
            ref = self.names[ref.lower()]
        m = _REF_RE.match(ref.replace(" ", "") if "'" not in ref else ref)
        if not m:
            raise Unsupported(f"reference {ref!r}")
        sheet = m.group("sheet")
        if sheet:
            sheet = sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet
        sheet = (sheet or cur_sheet).lower()
        ws = self.sheets.get(sheet)
        if ws is None:
            raise _Raise(XLError("#REF!"))
        a, b = m.group("a").replace("$", ""), (m.group("b") or m.group("a")).replace("$", "")

        def split(c, default_row):
            col = re.match(r"[A-Z]+", c, re.I).group(0)
            row = c[len(col):]
            return column_index_from_string(col.upper()), int(row) if row else default_row

        c1, r1 = split(a, 1)
        c2, r2 = split(b, ws.max_row)
        if not re.search(r"\d", a):  # whole column A:A
            r1, r2 = 1, ws.max_row
        r1, r2 = min(r1, r2), max(r1, r2)
        c1, c2 = min(c1, c2), max(c1, c2)
        return sheet, r1, r2, c1, c2

    def _ref(self, ref: str, cur_sheet: str):
        sheet, r1, r2, c1, c2 = self._range(ref, cur_sheet)
        if r1 == r2 and c1 == c2:
            return self.value(sheet, f"{get_column_letter(c1)}{r1}")
        return Array([[self.value(sheet, f"{get_column_letter(c)}{r}") for c in range(c1, c2 + 1)]
                      for r in range(r1, r2 + 1)])

    def _eval(self, node, sheet):
        kind = node[0]
        if kind == "lit":
            return node[1]
        if kind == "ref":
            for scope in reversed(self.scopes):  # LET-bound names shadow references
                if node[1].lower() in scope:
                    return scope[node[1].lower()]
            return self._ref(node[1], sheet)
        if kind == "neg":
            v = self._eval(node[1], sheet)
            return _elementwise(lambda x, _: -_num(x), v, 0.0) if isinstance(v, Array) else -_num(v)
        if kind == "pct":
            v = self._eval(node[1], sheet)
            return _elementwise(lambda x, _: _num(x) / 100, v, 0.0) if isinstance(v, Array) else _num(v) / 100
        if kind == "bin":
            a = self._eval(node[2], sheet)
            b = self._eval(node[3], sheet)
            return _elementwise(lambda x, y: _binop(node[1], x, y), a, b)
        if kind == "func":
            return self._call(node[1], node[2], sheet)
        raise Unsupported(kind)

    def _call(self, name, argn, sheet):
        ev = lambda n: self._eval(n, sheet)  # noqa: E731
        if name == "IF":
            cond = ev(argn[0])
            if isinstance(cond, Array):
                raise Unsupported("array IF")
            if _bool(cond):
                return ev(argn[1]) if len(argn) > 1 else True
            return ev(argn[2]) if len(argn) > 2 else False
        if name == "LET":
            if len(argn) < 3 or len(argn) % 2 == 0:
                raise _Raise(XLError("#VALUE!"))
            scope: dict = {}
            self.scopes.append(scope)
            try:
                for n, v in zip(argn[:-1:2], argn[1:-1:2]):
                    if n[0] != "ref":
                        raise Unsupported("LET name")
                    scope[n[1].lower()] = ev(v)
                return ev(argn[-1])
            finally:
                self.scopes.pop()
        if name in ("IFERROR", "IFNA"):
            try:
                v = ev(argn[0])
            except _Raise:
                return ev(argn[1])
            return ev(argn[1]) if isinstance(v, XLError) else v
        if name == "IFS":
            for c, v in zip(argn[::2], argn[1::2]):
                if _bool(ev(c)):
                    return ev(v)
            raise _Raise(XLError("#N/A"))
        args = [ev(a) for a in argn]
        fn = FUNCTIONS.get(name)
        if fn is None:
            raise Unsupported(f"function {name}")
        return fn(*args)


def _logical(args):
    out = []
    for a in args:
        if isinstance(a, Array):
            for v in a.flat():
                if isinstance(v, XLError):
                    raise _Raise(v)
                if isinstance(v, bool) or _is_num(v):
                    out.append(_bool(v))
        else:
            out.append(_bool(a))
    return out


def _avg(xs):
    if not xs:
        raise _Raise(XLError("#DIV/0!"))
    return sum(xs) / len(xs)


def _round(x, n=0.0, mode="half"):
    x, n = _num(x), int(_num(n))
    f = 10.0 ** n
    y = abs(x) * f
    if mode == "half":
        y = math.floor(y + 0.5 + 1e-9)
    elif mode == "up":
        y = math.ceil(y - 1e-9)
    else:
        y = math.floor(y + 1e-9)
    return math.copysign(y / f, x)


def _each(fn):
    """A one-value function given a range applies to each cell, as in Excel's array
    evaluation: SUMPRODUCT(ISNUMBER(F2:F41)*...) tests every cell, not the range."""
    def lifted(v, *rest):
        if isinstance(v, Array):
            return Array([[fn(x, *rest) for x in row] for row in v.rows])
        return fn(v, *rest)
    return lifted


def _sumproduct(*arrays):
    flats = [a.flat() if isinstance(a, Array) else [a] for a in arrays]
    n = len(flats[0])
    if any(len(f) != n for f in flats):
        raise _Raise(XLError("#VALUE!"))
    total = 0.0
    for i in range(n):
        p = 1.0
        for f in flats:
            v = f[i]
            if isinstance(v, XLError):
                raise _Raise(v)
            p *= float(v) if _is_num(v) else 0.0  # Excel: text and booleans count as 0
        total += p
    return total


def _as_rows(a):
    return a.rows if isinstance(a, Array) else [[a]]


def _same(v, key, wildcard=False) -> bool:
    """Exact-match equality used by MATCH/VLOOKUP/XLOOKUP (text is case-insensitive)."""
    if isinstance(key, str) and isinstance(v, str):
        if wildcard and any(c in key for c in "*?"):
            return fnmatch.fnmatchcase(v.lower(), key.lower())
        return v.lower() == key.lower()
    if _is_num(key) and _is_num(v):
        return float(v) == float(key)
    if isinstance(key, bool) and isinstance(v, bool):
        return v is key
    return False


def _approx_pos(values, key, descending=False):
    """Approximate match on sorted data: last position <= key (or >= when descending)."""
    best = None
    for i, v in enumerate(values):
        if v is None or isinstance(v, XLError) or (_is_num(v) != _is_num(key)):
            continue
        if (not descending and _compare("<=", v, key)) or (descending and _compare(">=", v, key)):
            best = i
        elif best is not None:
            break
    return best


def _index(arr, row, col=None):
    rows = _as_rows(arr)
    r = int(_num(row))
    c = None if col is None else int(_num(col))
    if c is None:
        if len(rows) == 1:          # a single row: the index runs along it
            r, c = 1, r
        elif len(rows[0]) == 1:
            c = 1
        else:
            if r < 1 or r > len(rows):
                raise _Raise(XLError("#REF!"))
            return Array([rows[r - 1]])
    if r == 0 and c >= 1:
        return Array([[row_[c - 1]] for row_ in rows])
    if c == 0 and r >= 1:
        return Array([rows[r - 1]])
    if r < 1 or c < 1 or r > len(rows) or c > len(rows[0]):
        raise _Raise(XLError("#REF!"))
    return rows[r - 1][c - 1]


def _match(key, arr, mtype=1.0):
    values = arr.flat() if isinstance(arr, Array) else [arr]
    mtype = int(_num(mtype)) if mtype is not None else 1
    if mtype == 0:
        pos = next((i for i, v in enumerate(values) if _same(v, key, wildcard=True)), None)
    else:
        pos = _approx_pos(values, key, descending=mtype < 0)
    if pos is None:
        raise _Raise(XLError("#N/A"))
    return float(pos + 1)


def _vlookup(key, table, col, approx=True):
    rows = _as_rows(table)
    c = int(_num(col))
    if c < 1 or c > len(rows[0]):
        raise _Raise(XLError("#REF!"))
    first = [r[0] for r in rows]
    if approx is None or _bool(approx):
        pos = _approx_pos(first, key)
    else:
        pos = next((i for i, v in enumerate(first) if _same(v, key, wildcard=True)), None)
    if pos is None:
        raise _Raise(XLError("#N/A"))
    return rows[pos][c - 1]


def _xlookup(key, look, ret, if_not_found=None, match_mode=0.0, search_mode=1.0):
    lrows, rrows = _as_rows(look), _as_rows(ret)
    by_col = len(lrows[0]) == 1          # lookup vector is a column: results are rows
    values = [r[0] for r in lrows] if by_col else lrows[0]
    mode = int(_num(match_mode)) if match_mode is not None else 0
    order = range(len(values)) if (search_mode is None or _num(search_mode) >= 0) else range(len(values) - 1, -1, -1)
    pos = None
    if mode in (0, 2):
        pos = next((i for i in order if _same(values[i], key, wildcard=(mode == 2))), None)
    elif mode in (-1, 1):
        exact = next((i for i in order if _same(values[i], key)), None)
        if exact is not None:
            pos = exact
        else:
            cands = [i for i in order if _is_num(values[i]) == _is_num(key) and values[i] is not None
                     and _compare("<" if mode == -1 else ">", values[i], key)]
            if cands:
                pick = max if mode == -1 else min
                pos = pick(cands, key=lambda i: values[i] if _is_num(key) else str(values[i]).lower())
    else:
        raise Unsupported(f"XLOOKUP match_mode {mode}")
    if pos is None:
        if if_not_found is not None:
            return if_not_found
        raise _Raise(XLError("#N/A"))
    if by_col:
        row = rrows[pos]
        return row[0] if len(row) == 1 else Array([row])
    col = [r[pos] for r in rrows]
    return col[0] if len(col) == 1 else Array([[v] for v in col])


def _sel(rng, mask):
    return [v for v, keep in zip(rng.flat(), mask) if keep and _is_num(v)]


FUNCTIONS = {
    "AND": lambda *a: all(_logical(a)),
    "OR": lambda *a: any(_logical(a)),
    "NOT": lambda a: not _bool(a),
    "TRUE": lambda: True,
    "FALSE": lambda: False,
    "SUM": lambda *a: sum(_vals(a)),
    "AVERAGE": lambda *a: _avg(_vals(a)),
    "MAX": lambda *a: max(_vals(a), default=0.0),
    "MIN": lambda *a: min(_vals(a), default=0.0),
    "COUNT": lambda *a: float(sum(1 for x in a for v in (x.flat() if isinstance(x, Array) else [x]) if _is_num(v))),
    "COUNTA": lambda *a: float(sum(1 for x in a for v in (x.flat() if isinstance(x, Array) else [x]) if v is not None and v != "")),
    "COUNTBLANK": lambda a: float(sum(1 for v in a.flat() if v is None or v == "")),
    "COUNTIF": lambda r, c: float(sum(1 for v in r.flat() if _criterion(c)(v))),
    "COUNTIFS": lambda *a: float(sum(_ifs_mask(a))),
    "SUMIF": lambda r, c, s=None: sum(float(v) for v, keep in zip((s or r).flat(), [_criterion(c)(x) for x in r.flat()]) if keep and _is_num(v)),
    "SUMIFS": lambda s, *a: sum(_sel(s, _ifs_mask(a))),
    "AVERAGEIF": lambda r, c, s=None: _avg([float(v) for v, keep in zip((s or r).flat(), [_criterion(c)(x) for x in r.flat()]) if keep and _is_num(v)]),
    "AVERAGEIFS": lambda s, *a: _avg(_sel(s, _ifs_mask(a))),
    "MAXIFS": lambda s, *a: max(_sel(s, _ifs_mask(a)), default=0.0),
    "MINIFS": lambda s, *a: min(_sel(s, _ifs_mask(a)), default=0.0),
    "SUMPRODUCT": _sumproduct,
    "INDEX": _index,
    "MATCH": _match,
    "VLOOKUP": _vlookup,
    "XLOOKUP": _xlookup,
    "ROUND": lambda x, n=0.0: _round(x, n, "half"),
    "ROUNDUP": lambda x, n=0.0: _round(x, n, "up"),
    "ROUNDDOWN": lambda x, n=0.0: _round(x, n, "down"),
    "ABS": _each(lambda x: abs(_num(x))),
    "ROWS": lambda a: float(a.shape[0] if isinstance(a, Array) else 1),
    "COLUMNS": lambda a: float(a.shape[1] if isinstance(a, Array) else 1),
    "LEN": _each(lambda s: float(len(_text(s)))),
    "UPPER": _each(lambda s: _text(s).upper()),
    "LOWER": _each(lambda s: _text(s).lower()),
    "TRIM": _each(lambda s: " ".join(_text(s).split())),
    "CONCATENATE": lambda *a: "".join(_text(x) for x in a),
    "CONCAT": lambda *a: "".join(_text(v) for x in a for v in (x.flat() if isinstance(x, Array) else [x])),
    "ISNUMBER": _each(_is_num),
    "ISBLANK": _each(lambda v: v is None),
    "ISERROR": _each(lambda v: isinstance(v, XLError)),
    "N": _each(lambda v: float(v) if _is_num(v) else (1.0 if v is True else 0.0)),
}
