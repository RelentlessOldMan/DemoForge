r"""Self-check for xlcalc.py (not run during verification).

    .venv\Scripts\python.exe checks\_selftest_xlcalc.py

Builds a small workbook in memory and checks one or more formulas per supported
function against values worked out by hand. Exits non-zero on any mismatch.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlcalc  # noqa: E402

wb = openpyxl.Workbook()
cr = wb.active
cr.title = "Criteria"
for row in [["Parameter", "Value"], ["Max position error", 0.5], ["Min pass rate", 0.95], ["Max runtime", 50]]:
    cr.append(row)
rn = wb.create_sheet("Run data")  # a sheet name that needs quoting
rn.append(["RunID", "Scenario", "Err", "Rt", "Result", "Note"])
data = [("VR-1", "cruise", 0.2, 30.0, "PASS", None), ("VR-2", "maneuver", 0.8, 40.0, "FAIL", "slow?"),
        ("VR-3", "turn", 0.5, 55.0, "FAIL", None), ("VR-4", "maneuver", 0.6, 35.0, "FAIL", None)]
for r in data:
    rn.append(list(r))
wb.defined_names["MaxErr"] = DefinedName("MaxErr", attr_text="Criteria!$B$2")

R = "'Run data'!"
CASES = [
    # aggregates
    (f"=SUM({R}C2:C5)", 2.1), (f"=SUM({R}C2:C5,1,TRUE)", 4.1),
    (f"=COUNT({R}A2:D5)", 8.0), (f"=COUNTA({R}F2:F5)", 1.0), (f"=COUNTA({R}A:A)", 5.0),
    (f"=AVERAGE({R}C2:C5)", 0.525), (f"=MIN({R}D2:D5)", 30.0), (f"=MAX({R}C:C)", 0.8),
    ("=ABS(-2.5)", 2.5), ("=ROUND(2/3,3)", 0.667),
    # conditional family
    (f'=COUNTIF({R}E2:E5,"pass")', 1.0), (f'=COUNTIF({R}C2:C5,">"&Criteria!B2)', 2.0),
    (f'=COUNTIF({R}B2:B5,"man*")', 2.0), (f'=COUNTIF({R}C2:C5,"<=0.5")', 2.0),
    (f'=COUNTIFS({R}B2:B5,"maneuver",{R}D2:D5,"<=50")', 2.0),
    (f'=SUMIF({R}B2:B5,"maneuver",{R}C2:C5)', 1.4), (f'=SUMIFS({R}C2:C5,{R}B2:B5,"<>maneuver",{R}D2:D5,"<50")', 0.2),
    (f'=AVERAGEIF({R}B2:B5,"maneuver",{R}C2:C5)', 0.7), (f'=AVERAGEIFS({R}D2:D5,{R}E2:E5,"FAIL")', 130 / 3),
    (f'=MAXIFS({R}C2:C5,{R}B2:B5,"maneuver")', 0.8), (f'=MINIFS({R}D2:D5,{R}E2:E5,"FAIL")', 35.0),
    (f'=SUMPRODUCT(({R}C2:C5<=MaxErr)*({R}D2:D5<=Criteria!B4))', 1.0),
    # logic and errors
    ("=IF(AND(1<2,NOT(FALSE)),\"Y\",\"N\")", "Y"), ("=OR(FALSE,0,1)", True), ("=NOT(1)", False),
    ("=IFERROR(1/0,\"div\")", "div"), ("=IFS(1>2,\"a\",2>1,\"b\")", "b"),
    (f'=IFNA(MATCH("none",{R}B2:B5,0),-1)', -1.0),
    (f"=ISNUMBER({R}C2)", True), (f"=ISNUMBER({R}B2)", False), (f"=ISBLANK({R}F2)", True), (f"=ISBLANK({R}F3)", False),
    # lookups
    (f"=INDEX({R}C2:C5,2)", 0.8), (f"=INDEX({R}A2:F5,3,4)", 55.0), (f"=SUM(INDEX({R}C2:D5,0,2))", 160.0),
    (f'=MATCH("turn",{R}B2:B5,0)', 3.0), (f"=MATCH(0.55,{R}C2:C3,1)", 1.0),
    (f'=INDEX({R}D2:D5,MATCH("VR-4",{R}A2:A5,0))', 35.0),
    ('=INDEX(Criteria!B:B,MATCH("Max runtime",Criteria!A:A,0))', 50.0),
    ('=VLOOKUP("Max position error",Criteria!A2:B4,2,FALSE)', 0.5),
    ('=VLOOKUP("min pass*",Criteria!$A$2:$B$4,2,0)', 0.95),
    (f'=VLOOKUP("VR-3",{R}A2:E5,5,FALSE)', "FAIL"),
    ('=XLOOKUP("Max runtime",Criteria!A2:A4,Criteria!B2:B4)', 50.0),
    ('=XLOOKUP("nope",Criteria!A2:A4,Criteria!B2:B4,"missing")', "missing"),
    (f'=XLOOKUP("maneuver",{R}B2:B5,{R}A2:A5,,0,-1)', "VR-4"),
    (f'=XLOOKUP(0.55,{R}C2:C5,{R}A2:A5,,1)', "VR-4"),
    # LET, also as Excel saves it
    ('=LET(lim,Criteria!B2,n,4,lim*n)', 2.0),
    ('=_xlfn.LET(_xlpm.x,Criteria!B4,_xlpm.x+1)', 51.0),
    # a Result-style formula and implicit error propagation
    (f'=IF(AND({R}C4<=MaxErr,{R}D4<=Criteria!$B$4),"PASS","FAIL")', "FAIL"),
    ("=1/0", xlcalc.XLError("#DIV/0!")),
]

ws = wb.create_sheet("T")
for i, (formula, want) in enumerate(CASES, start=1):
    ws[f"A{i}"] = formula
calc = xlcalc.Workbook(wb)
bad = 0
for i, (formula, want) in enumerate(CASES, start=1):
    try:
        got = calc.value("T", f"A{i}")
    except xlcalc.Unsupported as exc:
        got = f"Unsupported: {exc}"
    ok = (abs(got - want) < 1e-9) if isinstance(want, float) and isinstance(got, float) else got == want
    if not ok:
        bad += 1
        print(f"FAIL {formula}: got {got!r}, want {want!r}")
print(f"{len(CASES) - bad}/{len(CASES)} xlcalc cases pass")
sys.exit(1 if bad else 0)
